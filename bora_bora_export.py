"""
bora_bora_export.py
====================
Builds the LCOE verification workbook (Inputs / Tranche_Schedule / Trajectory /
Hybrid_LCOE) for a finished run, with LIVE Excel formulas, so the app's NPC and
LCOE can be audited and re-run in Excel (change an input and the sheet
recalculates).

Layout mirrors the hand-built verification workbook:
  Inputs            - discount / inflation / escalation switch, unit costs, lifetimes
  Tranche_Schedule  - nameplate capacity added per commissioning year and technology
  Trajectory        - annual energy figures read from the app's own run
  Hybrid_LCOE       - year-by-year CAPEX / O&M by technology, discounting, NPC, LCOE

Cost formulas read the Tranche_Schedule table with SUMIFS/COUNTIFS, so adding,
removing or moving a tranche there flows straight through (insert new rows
INSIDE the table, not below it, so the ranges grow). Replacement CAPEX is booked
in commissioning year + lifetime + 1, the same rule the engine uses.
"""

from io import BytesIO

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

TECH_ORDER = ["pv", "wind", "bess", "otec"]
TECH_LABEL = {"pv": "PV", "wind": "WIND", "bess": "BESS", "otec": "OTEC"}
CAP_UNIT = {"pv": "MWp", "wind": "MW", "bess": "MWh", "otec": "MW"}
CAPEX_UNIT = {"pv": "$/MWp", "wind": "$/MW", "bess": "$/MWh (installed)", "otec": "$/MW"}
OM_UNIT = {"pv": "$/MWp/yr", "wind": "$/MW/yr", "bess": "$/MWh/yr", "otec": "$/MW/yr"}

HEAD_FILL = PatternFill("solid", fgColor="1F4E79")
HEAD_FONT = Font(bold=True, color="FFFFFF")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
NOTE_FONT = Font(italic=True, color="555555")
BOLD = Font(bold=True)
NO_REPLACEMENT_YEARS = 999   # lifetime value meaning "never replaced inside the horizon"


def _head(ws, row, labels, col0=1):
    for i, text in enumerate(labels):
        c = ws.cell(row, col0 + i, text)
        c.fill, c.font = HEAD_FILL, HEAD_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def build_lcoe_workbook(schedule, traj_df, unit_costs, nominal_discount_rate, inflation_rate,
                        analysis_start_year, analysis_end_year, escalate_capex, lifetimes,
                        degradation, app_lcoe, app_npv_cost):
    """Returns the workbook as bytes.

    schedule      TrancheSchedule from the run
    traj_df       trajectory DataFrame from the run (needs the *_mwh columns used below)
    unit_costs    dict as used by the engine (pv_capex_per_mwp, ..., otec_capex_fixed, ...)
    lifetimes     {'pv': yrs|None, 'wind': ..., 'bess': ..., 'otec': ...}; None = no replacement
    degradation   {'pv': frac, 'wind': ..., 'bess': ..., 'otec': ...} (informational only)
    app_lcoe / app_npv_cost   the app's own results, shown next to the sheet's for a tie-out
    """
    # ------------------------------------------------------------------ which technologies exist
    tranches = []
    for tech in TECH_ORDER:
        for t in getattr(schedule, tech):
            cap = t.energy_mwh if tech == "bess" else t.capacity_mw
            tranches.append((int(t.commissioning_year), tech, float(cap)))
    tranches.sort(key=lambda x: (x[0], TECH_ORDER.index(x[1])))
    techs = [k for k in TECH_ORDER if k in ("pv", "bess") or any(tr[1] == k for tr in tranches)]
    n_tr = len(tranches)

    wb = openpyxl.Workbook()

    # ------------------------------------------------------------------ Inputs
    wi = wb.active
    wi.title = "Inputs"
    wi["A1"] = "Bora Bora - LCOE model inputs (generated from the app run)"
    wi["A1"].font = Font(bold=True, size=13)
    wi["A2"] = ("Yellow cells are inputs - change them to test sensitivities. Unit costs are in analysis-start-year "
                "dollars. Lifetime 999 = no replacement inside the horizon.")
    wi["A2"].font = NOTE_FONT
    rows = [("Analysis start year", analysis_start_year), ("Analysis end year", analysis_end_year),
            ("Nominal discount rate (%)", nominal_discount_rate * 100), ("Inflation rate (%)", inflation_rate * 100),
            ("Escalate CAPEX with inflation (1=yes, 0=no)", 1 if escalate_capex else 0)]
    for i, (label, val) in enumerate(rows):
        wi.cell(4 + i, 1, label)
        c = wi.cell(4 + i, 2, val)
        c.fill = INPUT_FILL
    ref = {"start": "Inputs!$B$4", "end": "Inputs!$B$5", "disc": "Inputs!$B$6", "infl": "Inputs!$B$7",
           "esc": "Inputs!$B$8"}

    row = 10
    for tech in techs:
        wi.cell(row, 1, f"{TECH_LABEL[tech]} properties").font = BOLD
        row += 1
        refs = {}
        wi.cell(row, 1, f"{TECH_LABEL[tech]} CAPEX ({CAPEX_UNIT[tech]})")
        c = wi.cell(row, 2, unit_costs[f"{tech}_capex_per_mw{'p' if tech == 'pv' else ('h' if tech == 'bess' else '')}"])
        c.fill, c.number_format = INPUT_FILL, "#,##0"
        refs["capex"] = f"Inputs!$B${row}"
        row += 1
        if tech == "otec":
            wi.cell(row, 1, "OTEC fixed CAPEX per tranche ($)")
            c = wi.cell(row, 2, unit_costs.get("otec_capex_fixed", 0.0))
            c.fill, c.number_format = INPUT_FILL, "#,##0"
            refs["fixed"] = f"Inputs!$B${row}"
            row += 1
        wi.cell(row, 1, f"{TECH_LABEL[tech]} O&M ({OM_UNIT[tech]})")
        c = wi.cell(row, 2, unit_costs[f"{tech}_om_per_mw{'p' if tech == 'pv' else ('h' if tech == 'bess' else '')}_yr"])
        c.fill, c.number_format = INPUT_FILL, "#,##0"
        refs["om"] = f"Inputs!$B${row}"
        row += 1
        wi.cell(row, 1, f"{TECH_LABEL[tech]} degradation rate (%/yr) - informational, used by the app's dispatch only")
        wi.cell(row, 2, (degradation.get(tech) or 0.0) * 100)
        row += 1
        wi.cell(row, 1, f"{TECH_LABEL[tech]} lifetime (years)")
        life = lifetimes.get(tech)
        c = wi.cell(row, 2, int(life) if life else NO_REPLACEMENT_YEARS)
        c.fill = INPUT_FILL
        refs["life"] = f"Inputs!$B${row}"
        refs["life_val"] = int(life) if life else NO_REPLACEMENT_YEARS
        row += 2
        ref[tech] = refs
    wi.column_dimensions["A"].width = 62
    wi.column_dimensions["B"].width = 16

    # ------------------------------------------------------------------ Tranche_Schedule
    wt = wb.create_sheet("Tranche_Schedule")
    wt["A1"] = "Capacity-expansion schedule (from the app run)"
    wt["A1"].font = Font(bold=True, size=13)
    wt["A2"] = ("Nameplate (undegraded) capacity added at each commissioning year. BESS is INSTALLED MWh (the dispatch "
                "uses a Min-Max SoC window of that). Costs in Hybrid_LCOE read this table - insert new rows inside it.")
    wt["A2"].font = NOTE_FONT
    _head(wt, 4, ["Commissioning Year", "Technology", "Capacity Added", "Unit"])
    for i, (yr, tech, cap) in enumerate(tranches):
        r = 5 + i
        wt.cell(r, 1, yr)
        wt.cell(r, 2, TECH_LABEL[tech])
        wt.cell(r, 3, cap)
        wt.cell(r, 4, CAP_UNIT[tech])
    last = 4 + n_tr
    YR = f"Tranche_Schedule!$A$5:$A${last}"
    TC = f"Tranche_Schedule!$B$5:$B${last}"
    CP = f"Tranche_Schedule!$C$5:$C${last}"
    for col, w in zip("ABCD", (20, 14, 16, 10)):
        wt.column_dimensions[col].width = w

    # ------------------------------------------------------------------ Trajectory
    years = [int(y) for y in traj_df["year"]]
    wj = wb.create_sheet("Trajectory")
    wj["A1"] = "Annual energy - the app's own run (values), costed energy derived by formula"
    wj["A1"].font = Font(bold=True, size=13)
    wj["A2"] = ("Costed_Energy = Total_Served - Rooftop_Served: energy the PROJECT delivered to load. Rooftop is privately "
                "owned (not in project CAPEX), so its delivered energy (direct + its pro-rata share of battery "
                "throughput) is excluded. Check column must read ~0 (formula vs the app's own project_served figure).")
    wj["A2"].font = NOTE_FONT
    _head(wj, 5, ["Year", "Total_Demand_MWh", "Total_Unmet_MWh", "Total_Served_MWh", "OTEC_Gen_MWh",
                  "Rooftop_Exported_Gen_MWh", "BESS_Discharge_MWh", "Rooftop_Served_MWh", "Costed_Energy_MWh",
                  "App_Project_Served_MWh", "Check (should be ~0)"])
    for i, (_, r) in enumerate(traj_df.iterrows()):
        x = 6 + i
        wj.cell(x, 1, int(r["year"]))
        wj.cell(x, 2, float(r["total_demand_mwh"]))
        wj.cell(x, 3, float(r["total_unmet_mwh"]))
        wj.cell(x, 4, f"=B{x}-C{x}")
        wj.cell(x, 5, float(r["otec_gen_mwh"]))
        wj.cell(x, 6, float(r["rooftop_gen_exported_mwh"]))
        wj.cell(x, 7, float(r["bess_discharge_mwh"]))
        wj.cell(x, 8, float(r["rooftop_served_mwh"]))
        wj.cell(x, 9, f"=D{x}-H{x}")
        wj.cell(x, 10, float(r["project_served_mwh"]))
        wj.cell(x, 11, f"=I{x}-J{x}")
        for col in range(2, 12):
            wj.cell(x, col).number_format = "#,##0.0"
    for col in range(1, 12):
        wj.column_dimensions[get_column_letter(col)].width = 17
    wj.row_dimensions[5].height = 32

    # ------------------------------------------------------------------ Hybrid_LCOE
    wh = wb.create_sheet("Hybrid_LCOE")
    wh["A1"] = "Bora Bora Hybrid LCOE - cash-flow method (live formulas)"
    wh["A1"].font = Font(bold=True, size=13)
    wh["A2"] = ("CAPEX is booked in each tranche's commissioning year and again in each replacement year (commissioning + "
                "lifetime + 1); CAPEX is escalated at the inflation rate when the Inputs switch is 1, O&M always is. "
                "Costs and project-served energy are discounted at the nominal rate to the start year (t=0, DF=1). "
                "LCOE = NPV(costs) / NPV(energy). No salvage value at the horizon end.")
    wh["A2"].font = NOTE_FONT
    capex_cols = {t: 5 + i for i, t in enumerate(techs)}
    om_cols = {t: 5 + len(techs) + i for i, t in enumerate(techs)}
    c_total = 5 + 2 * len(techs)
    c_disc, c_energy, c_denergy = c_total + 1, c_total + 2, c_total + 3
    headers = ["Year", "t", "DF", "Infl_Index"] + [f"CAPEX_{TECH_LABEL[t]}" for t in techs] \
        + [f"OM_{TECH_LABEL[t]}" for t in techs] + ["Total_Costs", "Disc_Costs", "Energy_MWh", "Disc_Energy"]
    _head(wh, 4, headers)

    n = len(years)
    first, lastrow = 5, 4 + n
    span = int(analysis_end_year) - int(analysis_start_year)

    def cap_sum(label, yr_expr):
        return f'SUMIFS({CP},{YR},{yr_expr},{TC},"{label}")'

    def cnt(label, yr_expr):
        return f'COUNTIFS({YR},{yr_expr},{TC},"{label}")'

    for i, yr in enumerate(years):
        r = first + i
        wh.cell(r, 1, yr)
        wh.cell(r, 2, f"=A{r}-{ref['start']}")
        wh.cell(r, 3, f"=1/POWER(1+{ref['disc']}/100,B{r})")
        wh.cell(r, 4, f"=POWER(1+{ref['infl']}/100,B{r})")
        for tech in techs:
            lab = TECH_LABEL[tech]
            life_ref = ref[tech]["life"]
            k_max = max(1, -(-span // (ref[tech]["life_val"] + 1)))   # ceil: replacement cycles that can fit
            yr_terms = [f"$A{r}"] + [f"$A{r}-{k}*({life_ref}+1)" for k in range(1, k_max + 1)]
            cap_terms = "+".join(cap_sum(lab, y) for y in yr_terms)
            expr = f"({cap_terms})*{ref[tech]['capex']}"
            if tech == "otec":
                cnt_terms = "+".join(cnt(lab, y) for y in yr_terms)
                expr += f"+({cnt_terms})*{ref[tech]['fixed']}"
            wh.cell(r, capex_cols[tech], f"=({expr})*IF({ref['esc']}=1,$D{r},1)")
            wh.cell(r, om_cols[tech],
                    f'=SUMIFS({CP},{YR},"<="&$A{r},{TC},"{lab}")*{ref[tech]["om"]}*$D{r}')
        first_c, last_c = get_column_letter(5), get_column_letter(c_total - 1)
        wh.cell(r, c_total, f"=SUM({first_c}{r}:{last_c}{r})")
        wh.cell(r, c_disc, f"={get_column_letter(c_total)}{r}*C{r}")
        wh.cell(r, c_energy, f"=Trajectory!I{6 + i}")
        wh.cell(r, c_denergy, f"={get_column_letter(c_energy)}{r}*C{r}")
        wh.cell(r, 3).number_format = "0.0000"
        wh.cell(r, 4).number_format = "0.0000"
        for col in range(5, c_denergy + 1):
            wh.cell(r, col).number_format = "#,##0"

    tr = lastrow + 2
    wh.cell(tr, 1, "TOTAL").font = BOLD
    for col in list(range(5, c_denergy + 1)):
        letter = get_column_letter(col)
        c = wh.cell(tr, col, f"=SUM({letter}{first}:{letter}{lastrow})")
        c.number_format, c.font = "#,##0", BOLD

    cdisc_l, cden_l = get_column_letter(c_disc), get_column_letter(c_denergy)
    sr = tr + 2
    wh.cell(sr, 1, "NPV Costs ($)")
    wh.cell(sr, 3, f"={cdisc_l}{tr}").number_format = "#,##0"
    wh.cell(sr + 1, 1, "NPV Energy - costed (MWh)")
    wh.cell(sr + 1, 3, f"={cden_l}{tr}").number_format = "#,##0"
    wh.cell(sr + 2, 1, "LCOE (USD/MWh)").font = BOLD
    c = wh.cell(sr + 2, 3, f'=IFERROR(C{sr}/C{sr + 1},"")')
    c.number_format, c.font = "#,##0.00", BOLD
    wh.cell(sr + 4, 1, "App result - NPV Costs ($)")
    wh.cell(sr + 4, 3, float(app_npv_cost)).number_format = "#,##0"
    wh.cell(sr + 5, 1, "App result - LCOE (USD/MWh)")
    wh.cell(sr + 5, 3, float(app_lcoe)).number_format = "#,##0.00"
    wh.cell(sr + 6, 1, "Sheet minus app (should be ~0 while the Inputs are unchanged)")
    wh.cell(sr + 6, 3, f"=C{sr}-C{sr + 4}").number_format = "#,##0.00"
    wh.cell(sr + 6, 4, f"=C{sr + 2}-C{sr + 5}").number_format = "0.0000"

    tb = sr + 8
    _head(wh, tb, ["Technology", "", "NPV CAPEX ($)", "NPV O&M ($)", "Total ($)", "LCOE Contribution ($/MWh)"])
    dfr = f"$C${first}:$C${lastrow}"
    for j, tech in enumerate(techs):
        r = tb + 1 + j
        cl, ol = get_column_letter(capex_cols[tech]), get_column_letter(om_cols[tech])
        wh.cell(r, 1, TECH_LABEL[tech])
        wh.cell(r, 3, f"=SUMPRODUCT({cl}{first}:{cl}{lastrow},{dfr})")
        wh.cell(r, 4, f"=SUMPRODUCT({ol}{first}:{ol}{lastrow},{dfr})")
        wh.cell(r, 5, f"=C{r}+D{r}")
        wh.cell(r, 6, f"=E{r}/$C${sr + 1}")
        for col in (3, 4, 5):
            wh.cell(r, col).number_format = "#,##0"
        wh.cell(r, 6).number_format = "#,##0.00"
    r = tb + 1 + len(techs)
    wh.cell(r, 1, "Total").font = BOLD
    for col in (3, 4, 5, 6):
        letter = get_column_letter(col)
        c = wh.cell(r, col, f"=SUM({letter}{tb + 1}:{letter}{r - 1})")
        c.font = BOLD
        c.number_format = "#,##0" if col < 6 else "#,##0.00"

    wh.column_dimensions["A"].width = 30
    for col in range(2, c_denergy + 1):
        wh.column_dimensions[get_column_letter(col)].width = 15
    wh.freeze_panes = "E5"
    wh.row_dimensions[4].height = 30

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
